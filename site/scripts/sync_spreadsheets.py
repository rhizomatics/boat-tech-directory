#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = ["numbers-parser>=4.19", "openpyxl>=3.1"]
# ///
"""Convert the source Apple Numbers spreadsheets into the committed .xlsx files.

The .xlsx files in site/spreadsheets/ are what the site build reads (see
build-spreadsheets.mjs), so CI never needs Numbers or numbers-parser.

Sources are listed in site/spreadsheets/sources.json. They live on the
maintainer's Mac, so a missing source is skipped quietly: other contributors and
CI are unaffected. A spreadsheet is only reconverted when the content of its
source changes (hash recorded in sources.lock.json), since .xlsx output isn't
byte-for-byte reproducible and would otherwise churn on every commit.

Usage:
  uv run site/scripts/sync_spreadsheets.py            # convert changed sources
  uv run site/scripts/sync_spreadsheets.py --force    # convert all available sources
  uv run site/scripts/sync_spreadsheets.py --stage    # also `git add` the results (pre-commit)
"""

import argparse
import hashlib
import json
import subprocess
import sys
from datetime import datetime, timedelta
from pathlib import Path

SHEETS_DIR = Path(__file__).resolve().parent.parent / "spreadsheets"
SOURCES = SHEETS_DIR / "sources.json"
LOCK = SHEETS_DIR / "sources.lock.json"

# Excel column width is measured in characters of the default font (~7 px each).
PX_PER_PT = 96 / 72
PX_PER_CHAR = 7


def log(msg):
    print(f"[sync-spreadsheets] {msg}", file=sys.stderr)


def source_hash(path: Path) -> str:
    """Hash the content of a .numbers file or package.

    For package (directory) documents only Index.zip and Data/ carry content;
    Metadata/ and previews change when the file is merely opened.
    """
    h = hashlib.sha256()
    if path.is_file():
        h.update(path.read_bytes())
        return h.hexdigest()
    parts = [path / "Index.zip", *sorted((path / "Data").rglob("*"))]
    for part in parts:
        if part.is_file():
            h.update(str(part.relative_to(path)).encode())
            h.update(part.read_bytes())
    return h.hexdigest()


def excel_value(value):
    if isinstance(value, timedelta):
        return value.total_seconds() / 86400  # Excel durations are fractions of a day
    if isinstance(value, datetime) and value.tzinfo:
        return value.replace(tzinfo=None)
    return value


def worksheet_title(sheet, table, used):
    title = sheet.name if len(sheet.tables) == 1 else f"{sheet.name} - {table.name}"
    for ch in "[]:*?/\\":
        title = title.replace(ch, " ")
    title = title.strip()[:31] or "Sheet"
    base, n = title, 2
    while title in used:
        suffix = f" {n}"
        title = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(title)
    return title


def convert(source: Path, target: Path):
    from numbers_parser import Document
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    doc = Document(str(source))
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    thin = Side(style="thin", color="BFBFBF")

    for sheet in doc.sheets:
        for table in sheet.tables:
            ws = wb.create_sheet(worksheet_title(sheet, table, used))
            # Table name is shown above the table in Numbers (e.g. the tidal planner's
            # "From: … To: … NM ->" line), so keep it as the centre page header.
            ws.oddHeader.center.text = " ".join(table.name.split())

            for r, row in enumerate(table.iter_rows(), start=1):
                for c, cell in enumerate(row, start=1):
                    value = excel_value(cell.value)
                    xc = ws.cell(row=r, column=c, value=value)
                    style = cell.style
                    if style is None:
                        continue
                    if style.bg_color and not isinstance(style.bg_color, list):
                        rgb = style.bg_color
                        xc.fill = PatternFill("solid", fgColor=f"{rgb.r:02X}{rgb.g:02X}{rgb.b:02X}")
                    if style.bold or r <= table.num_header_rows:
                        xc.font = Font(bold=True)
                    xc.alignment = Alignment(wrap_text=True, vertical="top")
                    xc.border = Border(left=thin, right=thin, top=thin, bottom=thin)

            for c in range(table.num_cols):
                ws.column_dimensions[get_column_letter(c + 1)].width = round(
                    table.col_width(c) * PX_PER_PT / PX_PER_CHAR, 1
                )
            for r in range(table.num_rows):
                ws.row_dimensions[r + 1].height = table.row_height(r)
            for merged in table.merge_ranges:
                ws.merge_cells(merged)
            if table.num_header_rows or table.num_header_cols:
                ws.freeze_panes = ws.cell(row=table.num_header_rows + 1, column=table.num_header_cols + 1)
            ws.sheet_properties.tabColor = None
            log(f"  {sheet.name} / {table.name.strip()[:40]}: {table.num_rows} rows × {table.num_cols} cols")

    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--force", action="store_true", help="convert even if the source is unchanged")
    parser.add_argument("--stage", action="store_true", help="git add converted files (for pre-commit)")
    args = parser.parse_args()

    sources = json.loads(SOURCES.read_text())
    lock = json.loads(LOCK.read_text()) if LOCK.exists() else {}
    changed = []

    for entry in sources:
        source = Path(entry["source"]).expanduser()
        target = SHEETS_DIR / entry["xlsx"]
        if not source.exists():
            log(f"skip {entry['xlsx']}: source not on this machine ({entry['source']})")
            continue
        digest = source_hash(source)
        if not args.force and target.exists() and lock.get(entry["xlsx"]) == digest:
            log(f"unchanged {entry['xlsx']}")
            continue
        log(f"converting {source.name} -> {entry['xlsx']}")
        try:
            convert(source, target)
        except Exception as e:  # keep the commit honest: a broken source should block it
            log(f"FAILED {source}: {e}")
            return 1
        lock[entry["xlsx"]] = digest
        changed.append(target)

    if changed:
        LOCK.write_text(json.dumps(lock, indent=2, sort_keys=True) + "\n")
        if args.stage:
            subprocess.run(["git", "add", *map(str, changed), str(LOCK)], check=True)
            log(f"staged {len(changed)} spreadsheet(s)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
